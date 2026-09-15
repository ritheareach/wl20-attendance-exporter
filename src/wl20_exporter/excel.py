"""Excel (.xlsx) and CSV export.

The workbook mirrors the layout of the office's FaceGO attendance log
(`.recording_log/<Month>_<Year>_attendance.xlsx`): one sheet per day named
DD-MM-YYYY, newest first, with the same seven columns — No., Staff ID,
Staff Name, First Check-in, Last Check-out, Total Hours, Status — the same
header style (white on #366092) and the same column widths. Times are HH:MM
and hours read "9h19mn", exactly like that log, so HR reads both files the
same way.

Optional extras (off by default) keep the tool useful for troubleshooting:
`include_punches` adds the raw punch list per day (FaceGO's "raw log"
equivalent) and `include_details` adds the Device Info and Diagnostics sheets.
"""

from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .models import DailyRow, DeviceRead, PunchRecord
from .summary import build_daily_rows, summary_totals

# Same palette as the FaceGO log and the other office reports.
HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=12, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=12, bold=True, color="366092")
DATA_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
WARN_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")

TIME_FMT = "hh:mm"
DATE_FMT = "dd-mm-yyyy"

# Column sets + widths copied from the FaceGO log.
DAY_HEADERS = ["No.", "Staff ID", "Staff Name", "First Check-in", "Last Check-out",
               "Total Hours", "Status"]
DAY_WIDTHS = [5, 13, 22, 16, 16, 13, 12]
PUNCH_HEADERS = ["No.", "Staff ID", "Staff Name", "Timestamp", "Punch"]
PUNCH_WIDTHS = [5, 13, 22, 21, 12]


def _style_header(sheet, row: int, headers: Sequence[str]) -> None:
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 26


def _autosize(sheet, widths: Sequence[int]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _setup_print(sheet, landscape: bool = True, repeat_header: bool = False) -> None:
    """Make the sheet print/export as one page wide instead of column fragments."""
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    if repeat_header:
        sheet.print_title_rows = "1:1"


def sheet_name(day: date) -> str:
    """Day sheets are named DD-MM-YYYY, like every other date we show."""
    return day.strftime("%d-%m-%Y")


def _write_day_sheet(sheet, rows: Sequence[DailyRow]) -> None:
    """One day of attendance — the same seven columns as the FaceGO log."""
    _style_header(sheet, 1, DAY_HEADERS)
    sheet.freeze_panes = "A2"
    for position, row in enumerate(rows, start=1):
        line = position + 1
        sheet.cell(row=line, column=1, value=position).font = DATA_FONT
        sheet.cell(row=line, column=2, value=row.user_id).font = DATA_FONT
        sheet.cell(row=line, column=3, value=row.name or "").font = DATA_FONT
        if row.first_in:
            first = sheet.cell(row=line, column=4, value=row.first_in.time())
            first.number_format = TIME_FMT
            first.font = DATA_FONT
        # Like the FaceGO log: a day with no check-out yet leaves the last
        # check-out and the total hours blank and reads "Present".
        if row.last_out and row.has_checkout:
            last = sheet.cell(row=line, column=5, value=row.last_out.time())
            last.number_format = TIME_FMT
            last.font = DATA_FONT
            sheet.cell(row=line, column=6, value=row.hours_text).font = DATA_FONT
        sheet.cell(row=line, column=7, value=row.status_text).font = DATA_FONT
        for column in (1, 4, 5, 6, 7):
            sheet.cell(row=line, column=column).alignment = CENTER
    _autosize(sheet, DAY_WIDTHS)
    _setup_print(sheet, landscape=True, repeat_header=True)


def _write_punch_sheet(sheet, records: Sequence[PunchRecord]) -> None:
    """Raw punch list for a day — the exporter's raw-log equivalent."""
    _style_header(sheet, 1, PUNCH_HEADERS)
    sheet.freeze_panes = "A2"
    for position, record in enumerate(records, start=1):
        line = position + 1
        sheet.cell(row=line, column=1, value=position).font = DATA_FONT
        sheet.cell(row=line, column=2, value=record.user_id).font = DATA_FONT
        sheet.cell(row=line, column=3, value=record.name or "").font = DATA_FONT
        stamp = sheet.cell(row=line, column=4, value=record.timestamp)
        stamp.number_format = "dd-mm-yyyy hh:mm:ss"
        stamp.font = DATA_FONT
        sheet.cell(row=line, column=5, value=record.punch_text).font = DATA_FONT
        sheet.cell(row=line, column=1).alignment = CENTER
    _autosize(sheet, PUNCH_WIDTHS)
    _setup_print(sheet, landscape=True, repeat_header=True)


def _write_info_sheet(sheet, read: DeviceRead, start: Optional[date], end: Optional[date],
                      record_count: int) -> None:
    info = read.info
    if start and end:
        span = f"{start:%d-%m-%Y} .. {end:%d-%m-%Y}"
    elif start:
        span = f"from {start:%d-%m-%Y}"
    elif end:
        span = f"until {end:%d-%m-%Y}"
    else:
        span = "everything on the terminal"
    pairs = [
        ("Report", "WL20 Attendance Exporter"),
        ("Exported at", datetime.now().strftime("%d-%m-%Y %H:%M:%S")),
        ("App version", __version__),
        ("Device name", info.name or "-"),
        ("Serial number", info.serial or "-"),
        ("Firmware", info.firmware or "-"),
        ("Address", f"{info.host}:{info.port}"),
        ("Users on terminal", len(read.users)),
        ("Records decoded", len(read.records)),
        ("Records exported", record_count),
        ("Read duration (s)", read.duration_seconds),
        ("Date range", span),
        ("Decoded format", read.report.format_label),
    ]
    for position, (key, value) in enumerate(pairs, start=1):
        sheet.cell(row=position, column=1, value=key).font = BOLD_FONT
        sheet.cell(row=position, column=2, value=value).font = DATA_FONT
    _autosize(sheet, [22, 46])
    _setup_print(sheet, landscape=False)


def _write_diagnostics_sheet(sheet, read: DeviceRead) -> None:
    report = read.report
    lines: List[tuple] = [
        ("Decoded format", report.format_label),
        ("Record size (bytes)", report.record_size),
        ("Payload offset", report.offset),
        ("Declared payload bytes", report.declared_bytes),
        ("Received payload bytes", report.received_bytes),
        ("Records before de-duplication", report.raw_records_read),
        ("Duplicate records dropped", report.duplicates_dropped),
        ("Users decoded", report.users_parsed),
        ("Used user fallback parser", "yes" if report.used_user_fallback else "no"),
        ("Used attendance fallback parser", "yes" if report.used_attendance_fallback else "no"),
        ("Device counter: users", read.info.users),
        ("Device counter: records", read.info.records),
        ("Device capacity: users", read.info.users_cap),
        ("Device capacity: records", read.info.records_cap),
        ("Read duration (s)", read.duration_seconds),
    ]
    for position, (key, value) in enumerate(lines, start=1):
        sheet.cell(row=position, column=1, value=key).font = BOLD_FONT
        sheet.cell(row=position, column=2, value=value).font = DATA_FONT

    row = len(lines) + 2
    sheet.cell(row=row, column=1, value="NOTES").font = TITLE_FONT
    for note in report.notes:
        row += 1
        sheet.cell(row=row, column=1, value=note).font = DATA_FONT
    row += 1
    sheet.cell(row=row, column=1, value="WARNINGS").font = TITLE_FONT
    if report.warnings:
        for warning in report.warnings:
            row += 1
            cell = sheet.cell(row=row, column=1, value=warning)
            cell.font = DATA_FONT
            cell.fill = WARN_FILL
    else:
        row += 1
        sheet.cell(row=row, column=1, value="none").font = DATA_FONT
    _autosize(sheet, [38, 70])
    _setup_print(sheet, landscape=False)


def export_workbook(path, read: DeviceRead, start: Optional[date] = None,
                    end: Optional[date] = None, include_punches: bool = False,
                    include_details: bool = False) -> Path:
    """Write the attendance workbook (one sheet per day) and return its path."""
    path = Path(path)
    records = list(read.records)
    rows = build_daily_rows(records)

    by_day: dict = {}
    for row in rows:
        by_day.setdefault(row.day, []).append(row)
    punches_by_day: dict = {}
    for record in records:
        punches_by_day.setdefault(record.day, []).append(record)

    workbook = Workbook()
    workbook.remove(workbook.active)

    for day in sorted(by_day, reverse=True):  # newest day first, like the FaceGO log
        day_rows = by_day[day]
        if include_punches:
            sheet = workbook.create_sheet(sheet_name(day) + " punches")
            _write_punch_sheet(sheet, punches_by_day.get(day, []))
        sheet = workbook.create_sheet(sheet_name(day))
        _write_day_sheet(sheet, day_rows)

    if not workbook.sheetnames:  # nothing in range: still produce a readable file
        sheet = workbook.create_sheet("Attendance")
        _style_header(sheet, 1, DAY_HEADERS)
        _autosize(sheet, DAY_WIDTHS)
        sheet.cell(row=2, column=1,
                   value="No attendance records in the selected range").font = DATA_FONT

    if include_details:
        _write_info_sheet(workbook.create_sheet("Device Info"), read, start, end, len(records))
        _write_diagnostics_sheet(workbook.create_sheet("Diagnostics"), read)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def export_csv(path, records: Iterable[PunchRecord]) -> Path:
    """Plain CSV copy of the raw records (UTF-8 with BOM so Excel opens it right)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["date", "time", "user_id", "name", "device_user_id",
                         "punch", "status", "uid", "timestamp"])
        for record in records:
            writer.writerow([
                record.timestamp.strftime("%d-%m-%Y"),
                record.timestamp.strftime("%H:%M:%S"),
                record.user_id,
                record.name,
                record.device_user_id,
                record.punch,
                record.status,
                record.uid,
                record.timestamp.strftime("%d-%m-%Y %H:%M:%S"),
            ])
    return path


def default_filename(start: Optional[date], end: Optional[date]) -> str:
    stamp = datetime.now().strftime("%H%M")
    if start and end:
        span = f"{start:%d-%m-%Y}_to_{end:%d-%m-%Y}"
    elif start:
        span = f"from_{start:%d-%m-%Y}"
    elif end:
        span = f"until_{end:%d-%m-%Y}"
    else:
        span = "all"
    return f"WL20_Attendance_{span}_{stamp}.xlsx"


def workbook_summary(rows: Sequence[DailyRow]) -> dict:
    """Small helper used by the CLI to report what was written."""
    totals = summary_totals(rows)
    return {"days": len({row.day for row in rows}), **totals}
