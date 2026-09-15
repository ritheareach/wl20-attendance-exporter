"""Excel/CSV export tests — the workbook must match the FaceGO log layout."""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, datetime, time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import load_workbook  # noqa: E402

from wl20_exporter import excel  # noqa: E402
from wl20_exporter.models import (  # noqa: E402
    DeviceInfo,
    DeviceRead,
    DeviceUser,
    ParseReport,
    PunchRecord,
)


def sample_read() -> DeviceRead:
    """Two days: a completed day (check-in + check-out) and an open one."""
    records = [
        PunchRecord(user_id="STF-0001", name="SOK Dara", uid=1,
                    timestamp=datetime(2026, 9, 14, 8, 1, 5), punch=0, status=1),
        PunchRecord(user_id="STF-0001", name="SOK Dara", uid=1,
                    timestamp=datetime(2026, 9, 14, 17, 30, 0), punch=1, status=1),
        PunchRecord(user_id="STF-0002", name="CHHIM Sokha", uid=2,
                    timestamp=datetime(2026, 9, 15, 8, 12, 30), punch=0, status=1),
    ]
    report = ParseReport(format_label="40-byte records", declared_bytes=120,
                         received_bytes=120, raw_records_read=3, users_parsed=2,
                         notes=["decoded range: 14-09-2026 08:01:05 .. 15-09-2026 08:12:30"])
    return DeviceRead(
        info=DeviceInfo(host="192.168.88.245", port=4370, name="WL20",
                        serial="A5KN203360148", firmware="Ver 6.60", users=2, records=3),
        users=[DeviceUser(uid=1, user_id="STF-0001", name="SOK Dara"),
               DeviceUser(uid=2, user_id="STF-0002", name="CHHIM Sokha")],
        records=records,
        report=report,
        duration_seconds=4.2,
    )


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "out.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_one_sheet_per_day_newest_first(self):
        excel.export_workbook(self.path, sample_read(), date(2026, 9, 14), date(2026, 9, 15))
        workbook = load_workbook(self.path)
        self.assertEqual(workbook.sheetnames, ["15-09-2026", "14-09-2026"])

    def test_columns_match_the_facego_log(self):
        excel.export_workbook(self.path, sample_read())
        sheet = load_workbook(self.path)["14-09-2026"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, ["No.", "Staff ID", "Staff Name", "First Check-in",
                                   "Last Check-out", "Total Hours", "Status"])
        row = [sheet.cell(row=2, column=column).value for column in range(1, 8)]
        self.assertEqual(row[0], 1)
        self.assertEqual(row[1], "STF-0001")
        self.assertEqual(row[2], "SOK Dara")
        self.assertEqual(row[3], time(8, 1, 5))
        self.assertEqual(row[4], time(17, 30, 0))
        self.assertEqual(row[5], "9h28mn")
        self.assertEqual(row[6], "Completed")
        self.assertEqual(sheet.cell(row=2, column=4).number_format, "hh:mm",
                         "times must read HH:MM like the office log")

    def test_open_day_stays_blank_and_reads_present(self):
        excel.export_workbook(self.path, sample_read())
        sheet = load_workbook(self.path)["15-09-2026"]
        row = [sheet.cell(row=2, column=column).value for column in range(1, 8)]
        self.assertEqual(row[3], time(8, 12, 30))
        self.assertIsNone(row[4], "no check-out yet")
        self.assertIsNone(row[5], "no hours to report yet")
        self.assertEqual(row[6], "Present")

    def test_header_style_matches_the_office_palette(self):
        excel.export_workbook(self.path, sample_read())
        cell = load_workbook(self.path)["14-09-2026"].cell(row=1, column=2)
        self.assertTrue(cell.font.bold)
        self.assertTrue(cell.font.color.rgb.endswith("FFFFFF"))
        self.assertTrue(cell.fill.fgColor.rgb.endswith("366092"))

    def test_punch_sheets_only_when_requested(self):
        excel.export_workbook(self.path, sample_read())
        self.assertEqual([name for name in load_workbook(self.path).sheetnames
                          if "punches" in name], [])

        excel.export_workbook(self.path, sample_read(), include_punches=True)
        workbook = load_workbook(self.path)
        self.assertEqual(workbook.sheetnames, ["15-09-2026 punches", "15-09-2026",
                                               "14-09-2026 punches", "14-09-2026"])
        sheet = workbook["14-09-2026 punches"]
        self.assertEqual([cell.value for cell in sheet[1]],
                         ["No.", "Staff ID", "Staff Name", "Timestamp", "Punch"])
        self.assertEqual(sheet.cell(row=2, column=4).value, datetime(2026, 9, 14, 8, 1, 5))
        self.assertEqual(sheet.cell(row=2, column=5).value, "Check In")

    def test_details_sheets_only_when_requested(self):
        excel.export_workbook(self.path, sample_read())
        self.assertNotIn("Device Info", load_workbook(self.path).sheetnames)

        excel.export_workbook(self.path, sample_read(), include_details=True)
        workbook = load_workbook(self.path)
        self.assertIn("Device Info", workbook.sheetnames)
        self.assertIn("Diagnostics", workbook.sheetnames)
        info_values = [cell.value for row in workbook["Device Info"].iter_rows() for cell in row]
        self.assertIn("A5KN203360148", info_values)
        diag = [cell.value for row in workbook["Diagnostics"].iter_rows() for cell in row]
        self.assertIn("40-byte records", diag)
        self.assertIn("none", diag)  # no warnings

    def test_empty_read_still_writes_a_readable_workbook(self):
        empty = DeviceRead(info=DeviceInfo(host="10.0.0.9"), report=ParseReport())
        excel.export_workbook(self.path, empty)
        sheet = load_workbook(self.path)["Attendance"]
        self.assertEqual([cell.value for cell in sheet[1]], excel.DAY_HEADERS)
        self.assertIn("No attendance records", sheet.cell(row=2, column=1).value)

    def test_sheet_names_use_dd_mm_yyyy(self):
        excel.export_workbook(self.path, sample_read())
        for name in load_workbook(self.path).sheetnames:
            self.assertRegex(name, r"^\d{2}-\d{2}-\d{4}$")


class CsvTests(unittest.TestCase):
    def test_csv_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "records.csv"
            excel.export_csv(path, sample_read().records)
            text = path.read_text(encoding="utf-8-sig")
            lines = text.strip().splitlines()
            self.assertEqual(len(lines), 4)
            self.assertTrue(lines[0].startswith("date,time,user_id"))
            self.assertIn("STF-0001", lines[1])
            self.assertIn("14-09-2026,08:01:05", lines[1])
            self.assertIn("14-09-2026 08:01:05", lines[1])


class FilenameTests(unittest.TestCase):
    def test_default_filename_uses_dd_mm_yyyy(self):
        name = excel.default_filename(date(2026, 9, 1), date(2026, 9, 15))
        self.assertTrue(name.startswith("WL20_Attendance_01-09-2026_to_15-09-2026_"))
        self.assertTrue(name.endswith(".xlsx"))
        self.assertIn("all", excel.default_filename(None, None))


if __name__ == "__main__":
    unittest.main()
